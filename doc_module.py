from docx import Document
from datetime import datetime

def create_document(content):
    doc = Document()
    doc.add_heading('AI Assistant Report', 0)
    doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(content)
    
    filename = f"AI_Answers_{datetime.now().strftime('%H%M%S')}.docx"
    doc.save(filename)
    return filename